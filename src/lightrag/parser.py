"""Untyped document parsing for the lightweight RAG core.

M1 supports PDF (pypdf), DOCX (python-docx) and TXT. Every block carries
`source_file`; PDF blocks also carry a 1-based `page` number. Errors are
structured and explicit so that the ingestion report never silently drops a
document.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

ERROR_UNSUPPORTED_FORMAT = "unsupported_format"
ERROR_EMPTY_DOCUMENT = "empty_document"
ERROR_CORRUPTED_FILE = "corrupted_file"
ERROR_FILE_NOT_FOUND = "file_not_found"
ERROR_PARSE_FAILURE = "parse_failure"


@dataclass(frozen=True)
class TextBlock:
    """One parsed text unit (a PDF page, DOCX paragraph/table, TXT paragraph)."""

    text: str
    source_file: str
    block_index: int
    page: int | None = None
    kind: str = "text"

    def to_dict(self) -> dict[str, Any]:
        return {
            "text": self.text,
            "source_file": self.source_file,
            "block_index": self.block_index,
            "page": self.page,
            "kind": self.kind,
        }


@dataclass
class ParsedDocument:
    """Uniform representation returned by every parser."""

    source_file: str
    file_type: str
    blocks: list[TextBlock] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(block.text for block in self.blocks)

    @property
    def char_count(self) -> int:
        return len(self.text)

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_file": self.source_file,
            "file_type": self.file_type,
            "block_count": len(self.blocks),
            "char_count": self.char_count,
            "metadata": self.metadata,
        }


class DocumentParseError(Exception):
    """Structured parse error: type + file name + reason."""

    def __init__(
        self,
        error_type: str,
        file_name: str,
        reason: str,
        *,
        source_file: str | None = None,
        details: dict[str, Any] | None = None,
    ) -> None:
        self.error_type = error_type
        self.file_name = file_name
        self.reason = reason
        self.source_file = source_file or file_name
        self.details = details or {}
        super().__init__(f"{error_type}: {file_name}: {reason}")

    def to_dict(self) -> dict[str, Any]:
        return {
            "type": self.error_type,
            "file": self.file_name,
            "source_file": self.source_file,
            "reason": self.reason,
            "details": self.details,
        }


def _source_file(path: Path) -> str:
    try:
        return path.resolve().relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def _check_path(path: Path) -> None:
    if not path.exists():
        raise DocumentParseError(
            ERROR_FILE_NOT_FOUND,
            path.name,
            "file does not exist",
            source_file=_source_file(path),
        )
    if not path.is_file():
        raise DocumentParseError(
            ERROR_PARSE_FAILURE,
            path.name,
            "path is not a regular file",
            source_file=_source_file(path),
        )
    if path.stat().st_size == 0:
        raise DocumentParseError(
            ERROR_EMPTY_DOCUMENT,
            path.name,
            "file is empty (0 bytes)",
            source_file=_source_file(path),
        )


def _parse_pdf(path: Path, source_file: str) -> ParsedDocument:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # pypdf raises several backend-specific errors
        raise DocumentParseError(
            ERROR_CORRUPTED_FILE,
            path.name,
            f"cannot open PDF: {exc}",
            source_file=source_file,
            details={"exception": type(exc).__name__},
        ) from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentParseError(
                ERROR_CORRUPTED_FILE,
                path.name,
                f"encrypted PDF cannot be opened: {exc}",
                source_file=source_file,
                details={"encrypted": True},
            ) from exc

    blocks: list[TextBlock] = []
    try:
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                blocks.append(
                    TextBlock(
                        text=text,
                        source_file=source_file,
                        block_index=len(blocks),
                        page=page_number,
                        kind="page",
                    )
                )
    except Exception as exc:
        raise DocumentParseError(
            ERROR_CORRUPTED_FILE,
            path.name,
            f"PDF page extraction failed: {exc}",
            source_file=source_file,
            details={"exception": type(exc).__name__},
        ) from exc

    if not blocks:
        raise DocumentParseError(
            ERROR_EMPTY_DOCUMENT,
            path.name,
            "PDF contains no extractable text (possibly a scanned/image-only PDF)",
            source_file=source_file,
            details={"pages": len(reader.pages)},
        )

    return ParsedDocument(
        source_file=source_file,
        file_type="pdf",
        blocks=blocks,
        metadata={"page_count": len(reader.pages), "parser": "pypdf"},
    )


def _iter_docx_blocks(document: Document) -> Iterable[Paragraph | Table]:
    """Yield paragraphs and tables in document order."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _table_text(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = " ".join(cell.text.split())
            if text:
                cells.append(text)
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows).strip()


def _parse_docx(path: Path, source_file: str) -> ParsedDocument:
    try:
        document = Document(str(path))
    except Exception as exc:
        raise DocumentParseError(
            ERROR_CORRUPTED_FILE,
            path.name,
            f"cannot open DOCX: {exc}",
            source_file=source_file,
            details={"exception": type(exc).__name__},
        ) from exc

    blocks: list[TextBlock] = []
    try:
        for item in _iter_docx_blocks(document):
            if isinstance(item, Paragraph):
                text = item.text.strip()
                kind = "paragraph"
            else:
                text = _table_text(item)
                kind = "table"
            if text:
                blocks.append(
                    TextBlock(
                        text=text,
                        source_file=source_file,
                        block_index=len(blocks),
                        page=None,
                        kind=kind,
                    )
                )
    except Exception as exc:
        raise DocumentParseError(
            ERROR_PARSE_FAILURE,
            path.name,
            f"DOCX block extraction failed: {exc}",
            source_file=source_file,
            details={"exception": type(exc).__name__},
        ) from exc

    if not blocks:
        raise DocumentParseError(
            ERROR_EMPTY_DOCUMENT,
            path.name,
            "DOCX contains no non-empty paragraphs or tables",
            source_file=source_file,
        )

    return ParsedDocument(
        source_file=source_file,
        file_type="docx",
        blocks=blocks,
        metadata={"block_count": len(blocks), "parser": "python-docx"},
    )


def _decode_txt(raw: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise DocumentParseError(
        ERROR_CORRUPTED_FILE,
        "",
        "TXT cannot be decoded as UTF-8 or GB18030",
    )


def _parse_txt(path: Path, source_file: str) -> ParsedDocument:
    try:
        text, encoding = _decode_txt(path.read_bytes())
    except DocumentParseError as exc:
        raise DocumentParseError(
            exc.error_type,
            path.name,
            exc.reason,
            source_file=source_file,
        ) from exc

    if not text.strip():
        raise DocumentParseError(
            ERROR_EMPTY_DOCUMENT,
            path.name,
            "TXT contains only whitespace",
            source_file=source_file,
        )

    # Blank lines separate logical paragraphs; if there are none, keep one block.
    paragraphs = [part.strip() for part in text.replace("\r\n", "\n").split("\n\n")]
    paragraphs = [part for part in paragraphs if part]
    if not paragraphs:
        paragraphs = [text.strip()]

    blocks = [
        TextBlock(
            text=part,
            source_file=source_file,
            block_index=index,
            page=None,
            kind="text",
        )
        for index, part in enumerate(paragraphs)
    ]
    return ParsedDocument(
        source_file=source_file,
        file_type="txt",
        blocks=blocks,
        metadata={"encoding": encoding, "parser": "builtin"},
    )


def parse_document(path: str | Path) -> ParsedDocument:
    """Parse one PDF/DOCX/TXT file into a `ParsedDocument`.

    Raises:
        DocumentParseError: with `error_type`, `file_name`, and `reason`.
    """
    file_path = Path(path)
    _check_path(file_path)
    source_file = _source_file(file_path)
    suffix = file_path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            ERROR_UNSUPPORTED_FORMAT,
            file_path.name,
            f"unsupported extension {suffix!r}; supported: {sorted(SUPPORTED_EXTENSIONS)}",
            source_file=source_file,
        )
    if suffix == ".pdf":
        return _parse_pdf(file_path, source_file)
    if suffix == ".docx":
        return _parse_docx(file_path, source_file)
    return _parse_txt(file_path, source_file)


def find_supported_documents(root: str | Path) -> list[Path]:
    """Return supported document paths below `root`, sorted for reproducibility."""
    root_path = Path(root)
    if not root_path.exists():
        raise FileNotFoundError(f"document directory does not exist: {root_path}")
    return sorted(
        (
            path
            for path in root_path.rglob("*")
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
        ),
        key=lambda item: item.as_posix(),
    )
